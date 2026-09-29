from pathlib import Path
import math
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT = Path('/Users/dannyo/Documents/CoffeeReferral')
WORK = ROOT / 'tmp/branding'
OUT = ROOT / 'output/branding'
WORK.mkdir(parents=True, exist_ok=True)
OUT.mkdir(parents=True, exist_ok=True)
FONTS = Path('/System/Library/Fonts/Supplemental')
C = dict(ink='#28211E', cream='#F7F2E8', sand='#E7DCCD', muted='#6B5E55',
         roast='#9F482E', bound='#285849', marked='#68405B', white='#FFFFFF')

def font(size, family='sans', bold=False):
    name = 'Georgia' if family == 'serif' else 'Arial'
    return ImageFont.truetype(str(FONTS / (name + (' Bold' if bold else '') + '.ttf')), size)

def panel(name, h=700, bg='cream'):
    im = Image.new('RGB', (1600, h), C[bg])
    return im, ImageDraw.Draw(im)

def label(d, xy, text, fill='muted', size=28):
    d.text(xy, text, font=font(size, bold=True), fill=C.get(fill, fill))

def text(d, xy, s, size=38, fill='ink', family='sans', bold=False):
    d.text(xy, s, font=font(size, family, bold), fill=C.get(fill, fill))

def save(im, name):
    p = WORK / (name + '.png')
    im.save(p)
    return p

def badge(d, xy, letters='RR', color='roast', size=120):
    x,y=xy
    d.rounded_rectangle((x,y,x+size,y+size), radius=int(size*.14), fill=C[color])
    f=font(int(size*.38),'serif',True)
    box=d.textbbox((0,0),letters,font=f)
    d.text((x+(size-(box[2]-box[0]))/2,y+(size-(box[3]-box[1]))/2-box[1]),letters,font=f,fill=C['cream'])

def button(d, xy, s, color='roast', width=345):
    x,y=xy
    d.rounded_rectangle((x,y,x+width,y+76), radius=12, fill=C[color])
    text(d,(x+25,y+20),s,30,'white',bold=True)

# A cover specimen shows the identity without making a claim about a live product.
im,d=panel('cover',950)
badge(d,(86,90),size=135)
text(d,(86,275),'ROAST',113,'ink','serif',True)
text(d,(86,405),'REGISTER',113,'ink','serif',True)
text(d,(91,567),'Find coffee that fits your taste.',43)
d.line((90,693,1507,693),fill=C['sand'],width=3)
text(d,(92,742),'BeanBound',62,'bound','serif',True)
text(d,(875,742),'BeanMarked',62,'marked','serif',True)
text(d,(95,831),'by RoastRegister',25,'muted')
text(d,(878,831),'by RoastRegister',25,'muted')
save(im,'cover')

im,d=panel('wordmarks',670)
text(d,(70,54),'ROAST REGISTER',84,'ink','serif',True)
label(d,(75,163),'PRIMARY HORIZONTAL WORDMARK',size=25)
d.line((75,228,1525,228),fill=C['sand'],width=2)
text(d,(75,278),'BeanBound',65,'bound','serif',True)
text(d,(79,365),'by RoastRegister',24,'muted')
text(d,(850,278),'BeanMarked',65,'marked','serif',True)
text(d,(854,365),'by RoastRegister',24,'muted')
badge(d,(78,468),size=115)
badge(d,(580,468),'BB','bound',115)
badge(d,(1080,468),'BM','marked',115)
label(d,(215,499),'DISCOVER',size=27)
label(d,(717,499),'TRAVEL',size=27)
label(d,(1217,499),'RATINGS',size=27)
save(im,'wordmarks')

im,d=panel('palette',650,'white')
swatches=[('ink','Espresso'),('cream','Cream'),('sand','Oat'),('muted','Stone'),
          ('roast','Terracotta'),('bound','Forest'),('marked','Plum'),('white','White')]
for i,(key,name) in enumerate(swatches):
    x=(i%4)*405; y=(i//4)*325
    d.rectangle((x,y,x+382,y+185),fill=C[key],outline=C['sand'],width=2)
    text(d,(x,y+205),name,35,bold=True)
    text(d,(x,y+254),C[key],29,'muted')
save(im,'palette')

im,d=panel('type',450)
label(d,(66,45),'GEORGIA  /  DISPLAY AND WORDMARKS',size=26)
text(d,(64,99),'A cup of your own.',95,'ink','serif')
label(d,(69,249),'ARIAL  /  INTERFACE AND BODY COPY',size=26)
text(d,(67,303),'Chocolate. Citrus. Something new.',46)
text(d,(69,371),'Light roast   /   Pour-over   /   Your tasting notes',31,'muted')
save(im,'type')

im,d=panel('icons',300)
# Shared 2 px at 24 px stroke language, drawn here at a larger teaching scale.
for idx,(key,title) in enumerate([('roast','Discover'),('bound','Travel'),('marked','Remember')]):
    x=100+idx*520; y=53; col=C[key]
    if idx==0:
        d.rounded_rectangle((x,y,x+105,y+130),radius=12,outline=col,width=8)
        for oy in (32,65,98):
            d.ellipse((x+18,y+oy-3,x+25,y+oy+4),fill=col)
            d.line((x+40,y+oy,x+83,y+oy),fill=col,width=8)
    elif idx==1:
        d.ellipse((x,y,x+128,y+128),outline=col,width=8)
        d.polygon([(x+87,y+27),(x+69,y+76),(x+27,y+100),(x+47,y+52)],outline=col,width=7)
    else:
        d.line([(x+16,y+126),(x+16,y),(x+108,y),(x+108,y+126),(x+62,y+94),(x+16,y+126)],fill=col,width=8)
    text(d,(x-15,218),title,36,bold=True)
save(im,'icons')

im,d=panel('applications',1120)
label(d,(55,42),'ROASTREGISTER  /  WEBSITE HERO',size=23)
text(d,(55,108),'Find coffee that',76,'ink','serif')
text(d,(55,197),'fits your taste.',76,'ink','serif')
text(d,(58,305),'Choose the flavors you love. Explore coffees',31)
text(d,(58,351),'from independent roasters that fit your cup.',31)
button(d,(58,418),'Find my coffee',width=300)
# A compact taste selector serves as the brand's graphic motif.
for i,(s,w) in enumerate([('Chocolate',215),('Citrus',155),('Floral',145)]):
    x=1050; y=121+i*108
    d.rounded_rectangle((x,y,x+w,y+66),radius=33,fill=C['white'],outline=C['roast'],width=3)
    text(d,(x+20,y+16),s,30)
d.line((57,568,1540,568),fill=C['sand'],width=3)
label(d,(58,618),'BEANBOUND',fill='bound',size=26)
text(d,(58,670),'Coffee worth',55,'ink','serif')
text(d,(58,740),'traveling for.',55,'ink','serif')
text(d,(58,830),'Find a stop worth the detour.',29)
button(d,(58,897),'Explore coffee stops','bound',370)
label(d,(842,618),'BEANMARKED',fill='marked',size=26)
text(d,(841,670),'Remember your',55,'ink','serif')
text(d,(841,740),'favorite cups.',55,'ink','serif')
text(d,(842,830),'Your taste. Your tasting history.',29)
button(d,(842,897),'Log a coffee','marked',310)
label(d,(60,1031),'by RoastRegister',size=23)
label(d,(844,1031),'by RoastRegister',size=23)
save(im,'applications')

def rgb(h): return RGBColor.from_string(h.lstrip('#'))
doc=Document()
s=doc.sections[0]
s.page_width=Inches(8.5); s.page_height=Inches(11)
s.top_margin=Inches(.64); s.bottom_margin=Inches(.64)
s.left_margin=s.right_margin=Inches(.75)
s.footer_distance=Inches(.30)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3','Caption']:
    st=doc.styles[name]
    st.font.name='Arial'
    st.font.bold=False
    st.font.italic=False
    st.font.color.rgb=RGBColor(0,0,0)
    st.paragraph_format.space_after=Pt(8)
    st.paragraph_format.line_spacing=1.12
    st.paragraph_format.left_indent=Pt(0)
    st.paragraph_format.right_indent=Pt(0)
    st.paragraph_format.first_line_indent=Pt(0)
doc.styles['Normal'].font.size=Pt(11)
for name,size in [('Title',28),('Heading 1',24),('Heading 2',14),('Heading 3',11)]:
    st=doc.styles[name]; st.font.name='Georgia' if name in ['Title','Heading 1'] else 'Arial'; st.font.size=Pt(size)
    st.font.bold=(name not in ['Title','Heading 1'])
    st.paragraph_format.space_before=Pt(13 if name!='Title' else 0)
    st.paragraph_format.space_after=Pt(8)
doc.styles['Subtitle'].font.size=Pt(13)
doc.styles['Caption'].font.size=Pt(9)
doc.styles['Caption'].font.italic=False
doc.styles['Caption'].font.color.rgb=rgb(C['muted'])
# Remove inherited theme fonts and paragraph rules from the base template.
for st in doc.styles:
    for node in st.element.xpath('.//w:pBdr'):
        node.getparent().remove(node)
    for node in st.element.xpath('.//w:rFonts'):
        for attr in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme','csTheme']:
            node.attrib.pop(qn('w:'+attr),None)
    for node in st.element.xpath('.//w:color'):
        for attr in ['themeColor','themeTint','themeShade']:
            node.attrib.pop(qn('w:'+attr),None)
# Quiet page numbering supports a reference manual.
foot=s.footer.paragraphs[0]
foot.alignment=WD_ALIGN_PARAGRAPH.RIGHT
r=foot.add_run('RoastRegister brand guide   |   ')
r.font.name='Arial';r.font.size=Pt(8);r.font.color.rgb=rgb(C['muted'])
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');foot._p.append(fld)
doc.core_properties.title='RoastRegister Brand Guide'
doc.core_properties.subject='Identity and usage standards for RoastRegister BeanBound and BeanMarked'
doc.core_properties.author='RoastRegister'
doc.core_properties.keywords='RoastRegister, BeanBound, BeanMarked, branding'

def p(t='',style=None,boldlead=None):
    el=doc.add_paragraph(style=style)
    if boldlead and t.startswith(boldlead):
        el.add_run(boldlead).bold=True;el.add_run(t[len(boldlead):])
    else: el.add_run(t)
    return el

def h(t,level=2): return doc.add_heading(t,level)
def page(n,title,deck):
    e=p(f'{n:02d}  /  BRAND STANDARDS');e.paragraph_format.space_after=Pt(8)
    e.paragraph_format.page_break_before=True
    for r in e.runs:r.font.size=Pt(9);r.font.bold=True
    h(title,1);p(deck)

def picture(name,caption=None,width=7):
    el=doc.add_paragraph();el.paragraph_format.space_after=Pt(8)
    shape=el.add_run().add_picture(str(WORK/(name+'.png')),width=Inches(width))
    shape._inline.docPr.set('descr',caption or name)
    if caption:p(caption,'Caption')

def table(headers,rows,widths):
    t=doc.add_table(rows=1, cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    for c,w in zip(t.columns,widths):c.width=Inches(w)
    for rowidx,vals in enumerate([headers]+rows):
        row=t.rows[0] if rowidx==0 else t.add_row()
        pr=row._tr.get_or_add_trPr()
        if rowidx==0:
            repeat=OxmlElement('w:tblHeader');pr.append(repeat)
        no=OxmlElement('w:cantSplit');pr.append(no)
        for j,(cell,txt) in enumerate(zip(row.cells,vals)):
            cell.width=Inches(widths[j]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cp=cell._tc.get_or_add_tcPr()
            shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'303030' if rowidx==0 else ('F5F5F3' if rowidx%2==0 else 'FFFFFF'));cp.append(shade)
            borders=OxmlElement('w:tcBorders')
            for side in ['top','left','bottom','right']:
                b=OxmlElement('w:'+side);b.set(qn('w:val'),'single');b.set(qn('w:sz'),'5');b.set(qn('w:color'),'D9D9D9');borders.append(b)
            cp.append(borders)
            margins=OxmlElement('w:tcMar')
            for side in ['top','left','bottom','right']:
                v=OxmlElement('w:'+side);v.set(qn('w:w'),'100');v.set(qn('w:type'),'dxa');margins.append(v)
            cp.append(margins)
            pp=cell.paragraphs[0];pp.paragraph_format.space_after=Pt(0);pp.paragraph_format.line_spacing=1.08
            rr=pp.add_run(txt);rr.font.name='Arial';rr.font.size=Pt(10)
            if rowidx==0:rr.bold=True;rr.font.color.rgb=rgb('#FFFFFF')
    e=p();e.paragraph_format.space_after=Pt(0);e.paragraph_format.line_spacing=Pt(8);e.paragraph_format.space_before=Pt(0)
    for r in e.runs:r.font.size=Pt(3)
    return t

def link(title,url):
    pp=p()
    hh=OxmlElement('w:hyperlink');hh.set(qn('r:id'),pp.part.relate_to(url,RT.HYPERLINK,is_external=True))
    rr=OxmlElement('w:r');pr=OxmlElement('w:rPr');color=OxmlElement('w:color');color.set(qn('w:val'),C['roast'][1:]);pr.append(color)
    un=OxmlElement('w:u');un.set(qn('w:val'),'single');pr.append(un);rr.append(pr)
    tt=OxmlElement('w:t');tt.text=title;rr.append(tt);hh.append(rr);pp._p.append(hh)

# 1 Cover
p('RoastRegister Brand Guide','Title')
p('Identity and usage standards for RoastRegister BeanBound and BeanMarked','Subtitle')
picture('cover','Shared wordmarks, warm neutrals and a distinct accent for each product.')
h('A guide for building the brand')
p('RoastRegister helps people discover coffee that fits their taste. BeanBound extends that curiosity into travel, and BeanMarked helps people remember what they have tried. This guide gives designers, writers and developers a shared foundation for presenting the family consistently.')
p('Use it for website design, product interfaces, social content and partner materials. Start with the brand purpose, then apply the naming, visual and writing rules to each new piece.')
p('Version 1.0  |  September 2026','Caption')

# 2 Foundation
page(2,'Brand foundation','Personal taste is the starting point. Discovery is the reason to return.')
h('Purpose')
p('Make specialty coffee easier to explore by connecting what a person enjoys with coffees, roasters and experiences worth trying.')
h('Positioning')
p('For curious coffee drinkers who want help choosing their next bag, RoastRegister is a guide to independent specialty roasters organized around personal taste. It turns flavor preferences, roast level and brewing style into useful recommendations with a clear explanation of the fit.')
h('The promise')
p('Find coffee that fits your taste.')
p('Every recommendation should help someone make a choice. Explain the flavors or preferences behind a match, allow different tastes to coexist, and make the next step obvious.')
h('Who we serve')
p('Our first audience is the home coffee drinker who knows some of what they like but does not want to research every roaster. Experienced enthusiasts should find enough detail to explore, while newcomers should feel comfortable starting with familiar notes such as chocolate, nuts or citrus.')
p('BeanBound serves people looking for coffee stops while traveling. BeanMarked serves people who want a useful record of coffees they have tasted. Both inherit the same welcoming approach.')
h('Personality')
table(['Quality','How it appears'],[
    ['Curious','Invite exploration through flavors, places and roasters.'],
    ['Informed','Explain the reason behind a suggestion in everyday language.'],
    ['Welcoming','Treat every roast preference and brewing method with respect.'],
    ['Grounded','Describe what is known; keep claims close to the evidence.']
],[1.5,5.5])
p('The desired feeling is a well-kept coffee field journal: useful, warm and personal. Avoid prestige language that makes a person feel they need expertise before they can begin.')

# 3 Architecture
page(3,'Brand family','Use RoastRegister as the parent identity and give each offshoot a clear job.')
table(['Brand','Role','Core action'],[
    ['RoastRegister','Personalized coffee and roaster discovery','Find my coffee'],
    ['BeanBound','Coffee travel and destination discovery','Explore coffee stops'],
    ['BeanMarked','Coffee ratings and tasting history','Log a coffee']
],[1.6,3.6,1.8])
h('Name each product consistently')
p('Write RoastRegister, BeanBound and BeanMarked as single words with the capitalization shown here. Do not use BeanThere, the earlier working name, in new material. In prose, use the product name rather than abbreviations such as RR or BM.')
p('The parent display wordmark may read ROAST REGISTER on one line or two. This visual treatment is an exception; page titles, accessible labels, metadata and prose retain RoastRegister.')
h('Make the relationship visible')
p('Use BeanBound by RoastRegister and BeanMarked by RoastRegister on the first prominent mention of either offshoot. In a logo lockup, place the endorsement on its own line below the product name. Repeat the parent relationship in an About section or footer, rather than on every button.')
p('When space is tight, preserve the product name and move the endorsement to a nearby readable location. The three brands share typography, neutral colors, spacing, and a consistent voice.')
h('Suggested navigation')
p('On the parent site, use Find coffee, Roasters, and How it works. Once the offshoots are available, introduce them as BeanBound Travel and BeanMarked Ratings in the product menu. Descriptive labels help visitors learn what the names mean.')
h('Current and future experiences')
p('The present website provides coffee matching and roaster discovery. BeanBound and BeanMarked are planned extensions. Their examples in this guide describe the intended identity; public copy should only offer features that are available. Use Coming soon when introducing an unfinished product.')

# 4 Wordmarks
page(4,'Wordmarks and identifiers','Use typography as the main signature. Keep the identity clear at small sizes.')
picture('wordmarks','Wordmark direction with endorsed offshoots and compact monogram identifiers.')
h('Construction')
p('Set the wordmarks in Georgia Bold. Keep the parent in capitals and the offshoots in their established mixed case. Set by RoastRegister in Arial Regular below the offshoot at roughly one third of its wordmark type size. Align the left edges; use normal letter spacing.')
table(['Rule','Standard'],[
    ['Clear space','Leave at least half the capital R height on every side.'],
    ['Minimum size','Parent horizontal mark: 180 px wide or 38 mm in print. Offshoot wordmark: 140 px wide or 30 mm.'],
    ['Endorsement','At least 12 px on screen or 8 pt in print; enlarge the lockup if needed.'],
    ['Compact identifier','Use RR, BB or BM at 32 px and above. Use a single R, B or M for a 16 px favicon.']
],[1.4,5.6])
p('Use Espresso on light backgrounds or Cream on dark backgrounds. Product accents are permitted for the matching offshoot. Keep proportions fixed, use a quiet background, and avoid outlines, shadows, gradients, rotation or decorative coffee imagery inside the lettering.')

# 5 Colors
page(5,'Color palette','Warm neutral surfaces unite the family. Product accents provide orientation.')
picture('palette','Digital palette in sRGB. Hex values are the source of truth.')
p('Use Cream for page backgrounds, White for reading surfaces, and Espresso for primary text. Stone supports secondary copy; Oat is a quiet divider or surface color. Terracotta belongs to RoastRegister, Forest to BeanBound, and Plum to BeanMarked. Use the active product accent for its main action.')

def luminance(h):
    v=[int(h[i:i+2],16)/255 for i in (1,3,5)]
    linear=[x/12.92 if x<=.04045 else ((x+.055)/1.055)**2.4 for x in v]
    return sum(a*b for a,b in zip(linear,[.2126,.7152,.0722]))
def contrast(a,b):
    v=sorted([luminance(C[a]),luminance(C[b])])
    return (v[1]+.05)/(v[0]+.05)
rows=[]
names=dict(ink='Espresso',muted='Stone',roast='Terracotta',bound='Forest',marked='Plum',white='White',cream='Cream')
for a,b,usage in [('ink','cream','Primary text'),('muted','cream','Secondary text'),('white','roast','RoastRegister buttons'),('white','bound','BeanBound buttons'),('white','marked','BeanMarked buttons')]:
    rows.append([f'{names[a]} on {names[b]}',f'{contrast(a,b):.2f}:1',usage])
table(['Text and background','Contrast','Use'],rows,[2.9,1.0,3.1])
p('These pairs exceed the 4.5:1 WCAG AA threshold for normal text at full opacity. Check each actual implementation, including hover, focus and disabled states. Oat is decorative: do not rely on it alone to identify an input boundary. Use Stone where a visible control edge is needed. [1, 2]','Caption')
p('Keep most of a composition neutral. Use accents sparingly, and pair color with text or an icon. Flavor-wheel colors describe flavors; they are separate from this brand palette. Prepare print colors for the chosen paper and printer, then approve a physical proof.')

# 6 Typography
page(6,'Typography and layout','Pair the character of Georgia with the clarity of Arial.')
picture('type','Georgia gives the brand an editorial voice; Arial keeps controls and data easy to scan.')
table(['Role','Typeface','Desktop / mobile'],[
    ['Wordmark','Georgia Bold','Scale the master wordmark'],
    ['Hero heading','Georgia Regular','56 / 36 px, 1.1 line height'],
    ['Section heading','Georgia Regular','36 / 28 px, 1.2 line height'],
    ['Card heading','Arial Bold','20 / 20 px, 1.3 line height'],
    ['Body text','Arial Regular','18 / 16 px, 1.5 line height'],
    ['Buttons and labels','Arial Bold','16 / 16 px, 1.3 line height'],
    ['Supporting metadata','Arial Regular','14 / 14 px, 1.4 line height']
],[1.6,1.9,3.5])
h('Spacing and composition')
p('Use an 8 px spacing rhythm, with 4 px for small adjustments. Set a 1200 px maximum content width, 24 px desktop gutters and 16 px mobile gutters. Give sections 64 to 96 px of vertical space on desktop and 40 to 56 px on mobile. Keep paragraphs around 55 to 75 characters per line.')
p('Use 8 px corner radii for buttons and 12 px for cards. Reserve pill shapes for filters and flavor chips. Avoid heavy shadows; use spacing and subtle surfaces to establish grouping. Make interactive targets at least 44 by 44 px as the brand design target.')
h('Type usage')
p('Use sentence case for interface labels and headings. Reserve all caps for the parent wordmark and short category labels. Use Georgia, Times New Roman, serif as the display fallback stack and Arial, Helvetica, sans-serif for body copy. Use licensed installations and do not bundle commercial font files without permission.')

# 7 Voice
page(7,'Voice and messaging','Sound like a helpful coffee enthusiast who listens before recommending.')
p('Speak directly to the reader using you and your. Be specific about flavor, roast and brewing style. Explain unfamiliar terms when they help someone choose, and let the product names carry most of the wordplay.')
table(['Use','Avoid'],[
    ['Choose the flavors you enjoy.','Unlock your ultimate sensory journey.'],
    ['A close match for your chocolate and nutty preferences.','Your scientifically perfect coffee.'],
    ['Like a fuller, roastier cup? Try these.','Real coffee lovers only drink light roast.'],
    ['No close matches yet. Try widening your roast preference.','No results. Your preferences are too restrictive.']
],[3.5,3.5])
h('Core messages')
p('RoastRegister: Find coffee that fits your taste.',boldlead='RoastRegister:')
p('Choose the flavors, roast level and brewing style you enjoy. Explore coffees from independent specialty roasters and see why each one fits your preferences.')
p('BeanBound: Coffee worth traveling for.',boldlead='BeanBound:')
p('Find coffee stops to explore on your next trip. Build a route around the cafés and roasters that catch your interest. Use this description when travel features are available.')
p('BeanMarked: Rate it. Remember it. Find it again.',boldlead='BeanMarked:')
p('Keep a record of the coffees you have tried, the way you brewed them and what you thought. Use your notes to choose what to revisit. Use this description when logging features are available.')
h('Trust in the details')
p('Use qualitative match language such as Close match and explain why. Do not present a preference score as a probability of enjoyment. Keep a user rating distinct from a recommendation, show the review count beside community averages, and treat an unrated coffee as Unrated rather than zero. Identify roaster tasting notes separately from a user\'s own observations.')

# 8 Imagery
page(8,'Photography and graphic style','Show the coffee, the people and the places with honest detail.')
el=doc.add_paragraph();el.alignment=WD_ALIGN_PARAGRAPH.CENTER;shape=el.add_run().add_picture(str(ROOT/'assets/coffee-roastery-hero.png'),width=Inches(6.1))
shape._inline.docPr.set('descr','Illustrative coffee roasting and tasting table from the existing site')
p('Existing illustrative website artwork. Use as a mood reference, not as a photograph of a named roaster or destination.','Caption')
h('Photography direction')
p('Favor natural light, believable color and tactile details: beans, paper labels, ceramic cups and working hands. Leave quiet space for headlines. RoastRegister focuses on coffee and tasting; BeanBound shows the surrounding street, café and route; BeanMarked shows the bag, brewed cup and notes.')
p('For named businesses or destinations, use authentic images with permission and retain required credits. Keep real product packaging legible and accurate. Avoid stock luxury cues, exaggerated steam, heavy orange filters and generated imagery presented as a real place.')
picture('icons','Use simple outline icons with the same stroke weight and level of detail.')
p('Draw interface icons on a 24 px grid with a 2 px stroke and rounded ends. Pair unfamiliar icons with labels. Use register, compass and bookmark motifs for discovery, travel and memory; keep them out of the wordmark itself. Charts and flavor wheels should retain readable labels and selected-state indicators beyond color alone.')

# 9 Applications
page(9,'Brand applications','Apply the same structure across products while changing the accent and the task.')
picture('applications','Illustrative interface compositions. BeanBound and BeanMarked show future product direction.')
h('Website and product surfaces')
p('Give each page one clear primary action. Lead with a useful promise, follow with a short explanation, and then offer the next step. Use product accents on controls; let coffee photography and tasting information provide the variety.')
h('Social and partner materials')
p('Use one headline, one image and a small wordmark. Keep text inside generous margins and check the mobile crop. For offshoot posts, include by RoastRegister on the opening image or in the profile description. A partner\'s logo should remain separate with its own clear space.')

# 10 Handoff
page(10,'Implementation and reference','Use this guide as the baseline for the first RoastRegister identity rollout.')
h('Apply the identity in order')
p('1. Establish the master wordmarks. Rebuild the specimens as vector artwork and inspect spacing at final sizes. Export full-color, one-color and reversed versions as SVG for digital use and vector PDF for print.')
p('2. Rebrand the current website. Update the visible name, page title, accessible labels and metadata. Introduce the core tagline, colors and typography together. Retain clear flavor selection, recommendation explanations and links to the roasters.')
p('3. Prepare compact assets. Create 16 and 32 px favicons, square social avatars, and reusable social layouts. Check the single-letter favicon at actual size. Keep endorsed offshoot wordmarks ready for their launch materials.')
p('4. Introduce the offshoots as they become usable. Add descriptive navigation and a shared About explanation. Describe account or profile continuity only when the products support it.')
h('Review before publishing')
p('Check spelling and capitalization, wordmark clear space, readable endorsements, text contrast, keyboard focus and mobile layout. Verify that copy describes available features, imagery represents the correct subject, and ratings are clearly distinguished from taste matches.')
h('Quick reference')
table(['Element','Default'],[
    ['Parent name','RoastRegister'],
    ['Offshoot names','BeanBound by RoastRegister; BeanMarked by RoastRegister'],
    ['Core promise','Find coffee that fits your taste.'],
    ['Type pairing','Georgia for display; Arial for interface and body'],
    ['Primary neutral pair','Espresso #28211E on Cream #F7F2E8'],
    ['Product accents','Terracotta #9F482E; Forest #285849; Plum #68405B']
],[1.6,5.4])
h('Accessibility references')
link('[1] W3C guidance on text contrast','https://www.w3.org/WAI/WCAG22/Understanding/contrast-minimum.html')
link('[2] W3C guidance on non-text contrast','https://www.w3.org/WAI/WCAG22/Understanding/non-text-contrast.html')
p('Contrast ratios in this guide were calculated from the listed sRGB hex values. Typography, layout and visual identity rules are RoastRegister design recommendations.','Caption')

doc.save(OUT/'RoastRegister-Brand-Guide.docx')
print(OUT/'RoastRegister-Brand-Guide.docx')
print('Contrast ratios:',rows)
